"""Build a tiny DOCX in-memory for golden coverage when python-docx is installed."""

from __future__ import annotations

import base64
from io import BytesIO

import pytest
from fastapi.testclient import TestClient

from app.main import app

client = TestClient(app)


def _make_docx_bytes() -> bytes | None:
    try:
        from docx import Document
    except ImportError:
        return None
    doc = Document()
    doc.add_heading("Etica no servico publico", level=1)
    doc.add_paragraph(
        "A etica no servico publico inclui deveres e vedacoes do servidor, "
        "transparencia, impessoalidade e responsabilidade perante a sociedade. "
        "Este capitulo integra material educacional de referencia para concursos "
        "publicos brasileiros e cobre os principios constitucionais da administracao."
    )
    doc.add_heading("Deveres", level=2)
    doc.add_paragraph(
        "Cumprir as leis, ser leal as instituicoes, prestar contas e zelar pelo "
        "patrimonio publico sao deveres permanentes do servidor. Observar as "
        "normas de conduta e o codigo de etica profissional quando aplicavel."
    )
    doc.add_heading("Principios", level=2)
    doc.add_paragraph(
        "Legalidade, impessoalidade, moralidade, publicidade e eficiencia formam "
        "o nucleo do art. 37 da Constituicao. A moralidade administrativa exige "
        "honestidade e boa-fe no exercicio da funcao publica."
    )
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_process_docx_when_python_docx_available():
    data = _make_docx_bytes()
    if data is None:
        pytest.skip("python-docx not installed")
    payload = {
        "documentId": "doc-docx-1",
        "contentType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "base64": base64.b64encode(data).decode("ascii"),
        "roleHint": "knowledge",
    }
    res = client.post("/process", json=payload)
    assert res.status_code == 200
    body = res.json()
    assert body["extractor"]["engine"] in ("python-docx", "docx-unavailable")
    assert body["stats"]["chars"] > 100
