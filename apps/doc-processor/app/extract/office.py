"""DOCX / DOC / PPTX extraction."""
from __future__ import annotations

import io
import subprocess
import tempfile
from pathlib import Path
from typing import Any

from ..schema import Block, Table


def extract(raw: bytes, fmt: str) -> tuple[list[Block], list[Table], dict[str, Any], dict[str, Any]]:
    if fmt == "doc":
        raw = _convert_with_libreoffice(raw, "docx")
        fmt = "docx"
        engine = "libreoffice+python-docx"
    elif fmt == "pptx":
        return _extract_pptx(raw)
    else:
        engine = "python-docx"

    try:
        from docx import Document  # type: ignore
    except ImportError as exc:
        raise ValueError("python_docx_unavailable") from exc

    document = Document(io.BytesIO(raw))
    blocks: list[Block] = []
    for p in document.paragraphs:
        text = (p.text or "").strip()
        if not text:
            continue
        style = (p.style.name if p.style is not None else "") or ""
        if style.lower().startswith("heading"):
            level = 1
            m = __import__("re").search(r"(\d+)", style)
            if m:
                level = int(m.group(1))
            blocks.append(Block(type="heading", text=text, level=level))
        else:
            blocks.append(Block(type="paragraph", text=text))

    tables: list[Table] = []
    for ti, table in enumerate(document.tables):
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        headers = rows[0] if rows else []
        body = rows[1:] if len(rows) > 1 else []
        md_lines = []
        if headers:
            md_lines.append("| " + " | ".join(headers) + " |")
            md_lines.append("| " + " | ".join("---" for _ in headers) + " |")
            for row in body:
                md_lines.append("| " + " | ".join(row) + " |")
        md = "\n".join(md_lines)
        tables.append(Table(id=f"t{ti}", headers=headers, rows=body, markdown=md))
        if md:
            blocks.append(Block(type="table", text=md))

    extractor = {"engine": engine, "version": "1", "options": {}}
    metadata = {"title": None, "author": None, "date": None, "sitename": None}
    return blocks, tables, extractor, metadata


def _extract_pptx(raw: bytes) -> tuple[list[Block], list[Table], dict[str, Any], dict[str, Any]]:
    try:
        from pptx import Presentation  # type: ignore
    except ImportError as exc:
        raise ValueError("python_pptx_unavailable") from exc
    prs = Presentation(io.BytesIO(raw))
    blocks: list[Block] = []
    for i, slide in enumerate(prs.slides):
        blocks.append(Block(type="heading", text=f"Slide {i + 1}", level=1))
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                blocks.append(Block(type="paragraph", text=shape.text.strip()))
    extractor = {"engine": "python-pptx", "version": "1", "options": {}}
    return blocks, [], extractor, {"title": None, "author": None, "date": None, "sitename": None}


def _convert_with_libreoffice(raw: bytes, target: str) -> bytes:
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "input.doc"
        src.write_bytes(raw)
        try:
            subprocess.run(
                ["libreoffice", "--headless", "--convert-to", target, "--outdir", tmp, str(src)],
                check=True,
                capture_output=True,
                timeout=120,
            )
        except Exception as exc:
            raise ValueError("libreoffice_convert_failed") from exc
        out = Path(tmp) / f"input.{target}"
        if not out.exists():
            raise ValueError("libreoffice_convert_missing_output")
        return out.read_bytes()
